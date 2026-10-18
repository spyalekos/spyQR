"""Scan history management with SecureStorage, disk fallback, and multi-format export (CSV, Excel, PDF)."""

import csv
import io
import json
import os
import platform
import subprocess
import sys
import time
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional
import flet as ft
import flet_secure_storage as fss


def get_downloads_dir() -> str:
    """Returns the user's Downloads directory across Windows, macOS, Linux, and Android."""
    # Android candidate paths
    android_candidates = [
        "/storage/emulated/0/Download",
        "/sdcard/Download",
    ]
    for candidate in android_candidates:
        try:
            if os.path.isdir(candidate) and os.access(candidate, os.W_OK):
                return candidate
        except Exception:
            pass

    # Desktop (Windows / macOS / Linux)
    desktop_dl = Path.home() / "Downloads"
    try:
        desktop_dl.mkdir(parents=True, exist_ok=True)
        return str(desktop_dl)
    except Exception:
        return str(Path.home())


def open_file_with_default_app(file_path: str, page: Optional[ft.Page] = None) -> bool:
    """Opens a file with the operating system's default application."""
    if not os.path.exists(file_path):
        return False
    try:
        sys_name = platform.system()
        if sys_name == "Windows":
            os.startfile(file_path)
            return True
        elif sys_name == "Darwin":
            subprocess.Popen(["open", file_path])
            return True
        elif sys_name == "Linux":
            if "ANDROID_ROOT" in os.environ or "ANDROID_DATA" in os.environ:
                if page:
                    page.launch_url(f"file://{file_path}")
                return True
            subprocess.Popen(["xdg-open", file_path])
            return True
        else:
            if page:
                page.launch_url(f"file://{file_path}")
            return True
    except Exception as ex:
        print(f"[open_file_with_default_app] Failed to open {file_path}: {ex}")
        return False


def get_pdf_font_path() -> Optional[str]:
    """Finds a usable TrueType Unicode font for PDF export with Greek language support."""
    current_dir = os.path.dirname(os.path.abspath(__file__))
    candidates = [
        os.path.join(current_dir, "..", "assets", "fonts", "DejaVuSans.ttf"),
        os.path.join(current_dir, "assets", "fonts", "DejaVuSans.ttf"),
        os.path.abspath(os.path.join("assets", "fonts", "DejaVuSans.ttf")),
    ]
    try:
        import barcode
        candidates.append(os.path.join(os.path.dirname(barcode.__file__), "fonts", "DejaVuSansMono.ttf"))
    except Exception:
        pass

    if sys.platform == "win32":
        candidates.extend([
            r"C:\Windows\Fonts\arial.ttf",
            r"C:\Windows\Fonts\calibri.ttf",
            r"C:\Windows\Fonts\segoeui.ttf",
        ])
    else:
        candidates.extend([
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "/usr/share/fonts/TTF/DejaVuSans.ttf",
            "/system/fonts/Roboto-Regular.ttf",
        ])

    for c in candidates:
        if c and os.path.isfile(c):
            return os.path.abspath(c)
    return None


def _resolve_fallback_file() -> str:
    """Returns a writable path for the history fallback JSON file on both Desktop and Android."""
    candidates = []
    env_data = os.environ.get("FLET_APP_STORAGE_DATA") or os.environ.get("SERIOUS_PYTHON_SITE_PACKAGES")
    if env_data:
        candidates.append(env_data)
    candidates.extend([
        "/data/user/0/gr.spyalekos.spyqr/files",
        "/data/data/gr.spyalekos.spyqr/files",
        os.path.expanduser("~"),
    ])
    for folder in candidates:
        try:
            if folder and os.path.isdir(folder) and os.access(folder, os.W_OK):
                return os.path.join(folder, ".spyqr_history.json")
        except Exception:
            continue
    return os.path.join(os.path.expanduser("~"), ".spyqr_history.json")


class HistoryManager:
    """Manages scan history persistence using flet-secure-storage with local disk fallback."""

    STORAGE_KEY = "spyqr_scan_history_v1"
    FALLBACK_FILE = _resolve_fallback_file()

    def __init__(self, page: ft.Page, secure_storage: Optional[fss.SecureStorage] = None):
        self.page = page
        self.secure_storage = secure_storage
        self.FALLBACK_FILE = _resolve_fallback_file()
        self.items: List[Dict[str, Any]] = []
        self._load_fallback()

    def _load_fallback(self) -> None:
        """Loads history from disk fallback if available, or seeds demo entries."""
        if os.path.exists(self.FALLBACK_FILE):
            try:
                with open(self.FALLBACK_FILE, "r", encoding="utf-8") as f:
                    data = json.load(f)
                    if isinstance(data, list) and len(data) > 0:
                        # Clean any legacy demo entries referencing flet.dev
                        cleaned = [
                            it for it in data
                            if "flet.dev" not in str(it.get("raw_data", ""))
                            and '"engine": "Flet' not in str(it.get("raw_data", ""))
                        ]
                        if cleaned:
                            self.items = cleaned
                            return
            except Exception as e:
                print(f"[HistoryManager] Fallback load error: {e}")

        now = time.time()
        self.items = [
            {
                "id": f"qr_{int(now * 1000)}_1",
                "timestamp_epoch": now - 120,
                "timestamp_str": time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(now - 120)),
                "qr_type": "PRODUCT",
                "raw_data": "5201054001050",
                "metadata": {"Κατηγορία": "Εμπορικό Προϊόν Λιανικής", "Πρότυπο Barcode": "EAN-13"},
            },
            {
                "id": f"qr_{int(now * 1000)}_2",
                "timestamp_epoch": now - 600,
                "timestamp_str": time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(now - 600)),
                "qr_type": "GEO",
                "raw_data": "geo:37.9755,23.7348",
                "metadata": {"latitude": 37.9755, "longitude": 23.7348},
            },
            {
                "id": f"qr_{int(now * 1000)}_3",
                "timestamp_epoch": now - 1800,
                "timestamp_str": time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(now - 1800)),
                "qr_type": "URL",
                "raw_data": "https://www.gov.gr",
                "metadata": {"url": "https://www.gov.gr"},
            },
        ]
        self._save_fallback()

    def _save_fallback(self) -> None:
        """Saves current items to disk fallback."""
        try:
            with open(self.FALLBACK_FILE, "w", encoding="utf-8") as f:
                json.dump(self.items, f, ensure_ascii=False, indent=2)
        except Exception as e:
            print(f"[HistoryManager] Fallback save error: {e}")

    async def load_history_async(self) -> List[Dict[str, Any]]:
        """Asynchronously loads history from SecureStorage."""
        if self.secure_storage:
            try:
                raw = await self.secure_storage.get(self.STORAGE_KEY)
                if raw:
                    data = json.loads(raw)
                    if isinstance(data, list):
                        self.items = data
                        return self.items
            except Exception as e:
                print(f"[HistoryManager] SecureStorage load exception: {e}")
        return self.items

    async def save_history_async(self) -> None:
        """Asynchronously saves history to SecureStorage and fallback."""
        self._save_fallback()
        if self.secure_storage:
            try:
                payload = json.dumps(self.items, ensure_ascii=False)
                await self.secure_storage.set(self.STORAGE_KEY, payload)
            except Exception as e:
                print(f"[HistoryManager] SecureStorage save exception: {e}")

    def add_entry(
        self,
        raw_data: str,
        qr_type: str = "TEXT",
        metadata: Optional[Dict[str, Any]] = None,
        on_saved: Optional[Callable[[], None]] = None,
    ) -> Dict[str, Any]:
        """Adds a newly scanned QR code entry to history."""
        # Avoid duplicate consecutive scans within 3 seconds
        current_time = time.time()
        if self.items:
            last = self.items[0]
            if last.get("raw_data") == raw_data and (current_time - last.get("timestamp_epoch", 0) < 3.0):
                return last

        entry = {
            "id": f"qr_{int(current_time * 1000)}",
            "timestamp_epoch": current_time,
            "timestamp_str": time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(current_time)),
            "qr_type": qr_type,
            "raw_data": raw_data,
            "metadata": metadata or {},
        }
        self.items.insert(0, entry)
        if len(self.items) > 500:
            self.items = self.items[:500]

        self._save_fallback()
        if self.secure_storage:
            try:
                self.page.run_task(self.save_history_async)
            except Exception:
                pass

        if on_saved:
            on_saved()
        return entry

    def delete_entry(self, entry_id: str) -> None:
        """Removes an entry by ID."""
        self.items = [item for item in self.items if item.get("id") != entry_id]
        self._save_fallback()
        if self.secure_storage:
            try:
                self.page.run_task(self.save_history_async)
            except Exception:
                pass

    def clear_all(self) -> None:
        """Clears all history entries."""
        self.items.clear()
        self._save_fallback()
        if self.secure_storage:
            try:
                self.page.run_task(self.save_history_async)
            except Exception:
                pass

    def get_stats(self) -> Dict[str, int]:
        """Computes count of items per QR type."""
        stats: Dict[str, int] = {}
        for item in self.items:
            t = item.get("qr_type", "TEXT")
            stats[t] = stats.get(t, 0) + 1
        return stats

    def export_csv(self, file_path: Optional[str] = None) -> str:
        """Exports history entries to a CSV file in the Downloads directory with UTF-8 BOM."""
        if not file_path:
            timestamp = time.strftime("%Y%m%d_%H%M%S")
            file_path = os.path.join(get_downloads_dir(), f"spyqr_history_{timestamp}.csv")

        sio = io.StringIO()
        writer = csv.writer(sio)
        writer.writerow(["A/A", "ID", "Ημερομηνία / Ώρα", "Τύπος", "Περιεχόμενο", "Λεπτομέρειες"])
        for idx, it in enumerate(self.items, 1):
            meta = it.get("metadata") or {}
            meta_str = " | ".join(f"{k}: {v}" for k, v in meta.items()) if isinstance(meta, dict) else str(meta)
            writer.writerow([
                idx,
                it.get("id", ""),
                it.get("timestamp_str", ""),
                it.get("qr_type", "TEXT"),
                it.get("raw_data", ""),
                meta_str,
            ])

        csv_bytes = ("\ufeff" + sio.getvalue()).encode("utf-8")
        os.makedirs(os.path.dirname(os.path.abspath(file_path)), exist_ok=True)
        with open(file_path, "wb") as f:
            f.write(csv_bytes)
        return file_path

    def export_excel(self, file_path: Optional[str] = None) -> str:
        """Exports history entries to an Excel (.xlsx) file in the Downloads directory with styling."""
        if not file_path:
            timestamp = time.strftime("%Y%m%d_%H%M%S")
            file_path = os.path.join(get_downloads_dir(), f"spyqr_history_{timestamp}.xlsx")

        import openpyxl
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Ιστορικό SpyQR"

        headers = ["A/A", "Ημερομηνία / Ώρα", "Τύπος", "Περιεχόμενο (Data)", "Λεπτομέρειες / Μεταδεδομένα"]
        ws.append(headers)

        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border
        ws.row_dimensions[1].height = 28

        for idx, item in enumerate(self.items, 1):
            meta = item.get("metadata") or {}
            meta_str = " | ".join(f"{k}: {v}" for k, v in meta.items()) if isinstance(meta, dict) else str(meta)
            row_data = [
                idx,
                item.get("timestamp_str", ""),
                item.get("qr_type", "TEXT"),
                item.get("raw_data", ""),
                meta_str,
            ]
            ws.append(row_data)
            row_idx = idx + 1
            ws.row_dimensions[row_idx].height = 22
            is_zebra = (idx % 2 == 0)

            for col_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = thin_border
                cell.font = Font(name="Calibri", size=10)
                if col_idx in (1, 2, 3):
                    cell.alignment = center_align
                else:
                    cell.alignment = left_align
                if is_zebra:
                    cell.fill = zebra_fill

        col_widths = {1: 8, 2: 22, 3: 16, 4: 55, 5: 45}
        for col_idx, width in col_widths.items():
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        os.makedirs(os.path.dirname(os.path.abspath(file_path)), exist_ok=True)
        wb.save(file_path)
        return file_path

    def export_pdf(self, file_path: Optional[str] = None) -> str:
        """Exports history entries to a landscape A4 PDF report in the Downloads directory."""
        if not file_path:
            timestamp = time.strftime("%Y%m%d_%H%M%S")
            file_path = os.path.join(get_downloads_dir(), f"spyqr_history_{timestamp}.pdf")

        from fpdf import FPDF

        font_path = get_pdf_font_path()
        use_custom_font = font_path is not None and os.path.exists(font_path)
        font_family = "DejaVu" if use_custom_font else "Helvetica"
        now_str = time.strftime("%Y-%m-%d %H:%M:%S")
        total_items = len(self.items)

        class HistoryPDF(FPDF):
            def header(self):
                self.set_font(font_family, "B" if use_custom_font else "B", 13)
                self.cell(0, 8, "SpyQR - Ιστορικό Σαρώσεων", align="C", new_x="LMARGIN", new_y="NEXT")
                self.set_font(font_family, "", 9)
                self.set_text_color(100, 116, 139)
                self.cell(
                    0, 6,
                    f"Ημερομηνία Εξαγωγής: {now_str}  |  Σύνολο Εγγραφών: {total_items}",
                    align="C", new_x="LMARGIN", new_y="NEXT",
                )
                self.set_text_color(0, 0, 0)
                self.ln(3)

            def footer(self):
                self.set_y(-12)
                self.set_font(font_family, "", 8)
                self.set_text_color(148, 163, 184)
                self.cell(0, 8, f"Σελίδα {self.page_no()}/{{nb}}  •  SpyQR App", align="C")

        pdf = HistoryPDF(orientation="L", unit="mm", format="A4")
        pdf.set_auto_page_break(auto=True, margin=15)

        if use_custom_font:
            pdf.add_font(font_family, "", font_path)
            pdf.add_font(font_family, "B", font_path)
            pdf.add_font(font_family, "I", font_path)

        pdf.add_page()
        pdf.set_font(font_family, "", 9)

        headers = ["A/A", "Ημερομηνία / Ώρα", "Τύπος", "Περιεχόμενο", "Λεπτομέρειες"]
        col_widths = (12, 42, 28, 130, 65)
        text_alignments = ("CENTER", "CENTER", "CENTER", "LEFT", "LEFT")

        with pdf.table(
            col_widths=col_widths,
            text_align=text_alignments,
            borders_layout="ALL",
            line_height=5.5,
        ) as table:
            h_row = table.row()
            for h in headers:
                h_row.cell(h)

            for idx, it in enumerate(self.items, 1):
                meta = it.get("metadata") or {}
                meta_str = " | ".join(f"{k}: {v}" for k, v in meta.items()) if isinstance(meta, dict) else str(meta)
                raw_data = str(it.get("raw_data", ""))
                # Sanitize high unicode/emojis (> 0xFFFF) for standard TTF BMP fonts
                raw_data = "".join(ch for ch in raw_data if ord(ch) < 0x10000)
                meta_str = "".join(ch for ch in meta_str if ord(ch) < 0x10000)
                if len(raw_data) > 300:
                    raw_data = raw_data[:297] + "..."
                if len(meta_str) > 200:
                    meta_str = meta_str[:197] + "..."

                row = table.row()
                row.cell(str(idx))
                row.cell(str(it.get("timestamp_str", "")))
                row.cell(str(it.get("qr_type", "TEXT")))
                row.cell(raw_data)
                row.cell(meta_str)

        os.makedirs(os.path.dirname(os.path.abspath(file_path)), exist_ok=True)
        pdf.output(file_path)
        return file_path

